"""解析测试：Docling 优先/降级、markdown/plain 定位、PDF/DOCX fallback、扫描件标记。"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.ingest import parsers
from app.ingest.parsers import docling_available, parse_document, parse_text_document


def _write_pdf(path: Path, text: str | None = "hello trade practice 2026") -> None:
    pdf = canvas.Canvas(str(path), pagesize=A4)
    if text:
        pdf.setFont("Helvetica", 12)
        pdf.drawString(72, 760, text)
    pdf.showPage()
    pdf.save()


def test_markdown_parse_headings_tables_and_line_locators(tmp_path: Path):
    path = tmp_path / "doc.md"
    path.write_text(
        "# 第一章 培养方案\n\n本专业培养方案说明段落。\n\n"
        "## 1.1 课程结构\n\n| 课程 | 学分 |\n| --- | --- |\n| 国际贸易实务 | 3 |\n",
        encoding="utf-8",
    )
    result = parse_text_document(path)
    assert result.extraction_status == "parsed"
    headings = [section.text for section in result.sections if section.block_type == "heading"]
    assert "第一章 培养方案" in headings
    tables = [section for section in result.sections if section.block_type == "table"]
    assert len(tables) == 1
    assert tables[0].locator["kind"] == "line_range"
    # Markdown 没有页码，不能伪造 page locator
    for section in result.sections:
        assert "page" not in section.locator
    assert result.sections[-1].heading_path == "第一章 培养方案 / 1.1 课程结构"


def test_plain_text_parse_line_ranges(tmp_path: Path):
    path = tmp_path / "doc.txt"
    path.write_text("第一段文字。\n\n第二段文字，包含更多内容。\n", encoding="utf-8")
    result = parse_text_document(path)
    assert len(result.sections) == 2
    assert result.sections[0].locator["kind"] == "line_range"
    assert result.sections[0].locator["line_start"] == 1
    assert result.sections[1].locator["line_start"] == 3


def test_docling_is_preferred_when_available(tmp_path: Path, monkeypatch):
    path = tmp_path / "doc.pdf"
    _write_pdf(path)
    called: list[bool] = []

    def fake_docling(path_arg, *, is_pdf):  # noqa: ANN001
        called.append(is_pdf)
        from app.ingest.parsers import ParseResult, ParsedSection

        return ParseResult(
            sections=[ParsedSection(text="docling 解析结果", locator={"kind": "page", "page": 1})],
            parser_version="docling-fake",
            extraction_status="parsed",
            docling_available=True,
        )

    monkeypatch.setattr(parsers, "docling_available", lambda: True)
    monkeypatch.setattr(parsers, "_parse_with_docling", fake_docling)
    result = parse_document(path, mime_type="application/pdf")
    assert called == [True]
    assert result.parser_version == "docling-fake"


def test_pypdf_fallback_records_parser_version_and_real_pages(tmp_path: Path, monkeypatch):
    path = tmp_path / "doc.pdf"
    _write_pdf(path, "hello trade practice 2026")
    monkeypatch.setattr(parsers, "docling_available", lambda: False)
    result = parse_document(path, mime_type="application/pdf")
    assert "fallback" in result.parser_version
    assert result.used_fallback is True
    assert result.pages_total == 1
    assert result.sections, "pypdf 应能提取 reportlab 生成的 ASCII 文本"
    assert result.sections[0].locator["kind"] == "page"
    assert result.sections[0].locator["page"] == 1


def test_scanned_pdf_is_marked_ocr_not_faked(tmp_path: Path, monkeypatch):
    path = tmp_path / "scan.pdf"
    _write_pdf(path, text=None)
    monkeypatch.setattr(parsers, "docling_available", lambda: False)
    result = parse_document(path, mime_type="application/pdf")
    assert result.extraction_status == "ocr"
    assert result.sections == []
    assert result.pages_empty == (1,)
    assert "scanned_or_empty_pages" in result.warnings


def test_docx_fallback_uses_block_locator_without_fake_pages(tmp_path: Path, monkeypatch):
    path = tmp_path / "doc.docx"
    document = Document()
    document.add_heading("第一章 培养方案", level=1)
    document.add_paragraph("本专业培养方案说明段落。")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "课程"
    table.cell(0, 1).text = "学分"
    table.cell(1, 0).text = "国际贸易实务"
    table.cell(1, 1).text = "3"
    document.save(str(path))
    monkeypatch.setattr(parsers, "docling_available", lambda: False)
    result = parse_document(path, mime_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert "fallback" in result.parser_version
    assert result.extraction_status == "partial"
    assert result.used_fallback is True
    assert any(section.block_type == "table" for section in result.sections)
    for section in result.sections:
        assert section.locator["kind"] == "docx_block"
        assert "page" not in section.locator
